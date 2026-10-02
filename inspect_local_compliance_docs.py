import json
import sys
from pathlib import Path

from docx import Document
from pypdf import PdfReader


for raw in sys.argv[1:]:
    path = Path(raw)
    print(f"\n===== {path} =====")
    if path.suffix.lower() == ".docx":
        document = Document(path)
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                print(f"[{paragraph.style.name}] {text}")
        for number, table in enumerate(document.tables, start=1):
            print(f"[TABLE {number}]")
            for row in table.rows:
                print(" | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))
    elif path.suffix.lower() == ".pdf":
        reader = PdfReader(path)
        for number, page in enumerate(reader.pages, start=1):
            print(f"[PAGE {number}]")
            print(page.extract_text() or "")
    else:
        print("Unsupported file type")
